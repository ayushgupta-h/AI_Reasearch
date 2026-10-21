"""
Report Exporter
===============
Converts a Markdown research report into PDF, DOCX, Markdown, and TXT formats
with support for citation styling (APA, MLA, Harvard), structured tables,
and rich styling.
"""

import re
import sys
from pathlib import Path
from typing import Dict, Optional, Tuple, List
from urllib.parse import urlparse

from rich.console import Console

if sys.platform.startswith("win"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

console = Console()


# ─── Citation Styling ─────────────────────────────────────────────────────────

def apply_citation_style(markdown_text: str, style: str = "apa") -> str:
    """
    Apply APA, MLA, or Harvard citation styling to references in the research report.
    Never invents missing authors or publication details.
    """
    style = (style or "apa").lower().strip()
    if style not in ("apa", "mla", "harvard"):
        style = "apa"

    lines = markdown_text.splitlines()
    output_lines = []
    in_ref_section = False
    has_existing_section = False

    for l in lines:
        if re.match(r"^#{1,3}\s+.*(?:References|Sources|Citations|Bibliography)", l, re.IGNORECASE):
            has_existing_section = True
            break

    for line in lines:
        if re.match(r"^#{1,3}\s+.*(?:References|Sources|Citations|Bibliography)", line, re.IGNORECASE):
            in_ref_section = True
            output_lines.append(line)
            output_lines.append(f"\n*Citations formatted in {style.upper()} style.*\n")
            continue
        elif in_ref_section and re.match(r"^#{1,3}\s+", line):
            in_ref_section = False

        if in_ref_section:
            # Match list item with link: * [text](url) — notes or * [text](url)
            m = re.match(r"^\s*([*\-+]|\d+\.)\s+\[([^\]]+)\]\((https?://[^\s)]+)\)(?:\s*[-—–]\s*(.*))?$", line)
            if m:
                bullet = m.group(1)
                text = m.group(2).strip()
                url = m.group(3).strip()
                desc = (m.group(4) or "").strip()
                domain = urlparse(url).netloc.replace("www.", "")

                # Extract year if present in parentheses or text
                year_match = re.search(r"\b(19\d\d|20\d\d)\b", f"{text} {desc}")
                year = year_match.group(1) if year_match else "n.d."

                # Clean source / author and title
                source = re.sub(r"\s*\(\d{4}\)", "", text).strip()
                clean_desc = re.sub(r"^\*+|\*+$", "", desc).strip().rstrip(".")
                title = clean_desc if clean_desc else source
                author = source if source != title else domain

                if style == "mla":
                    formatted = f'"{title}." *{domain}*, {year}, {url}.'
                elif style == "harvard":
                    formatted = f'{author} ({year}) *{title}*. Available at: <{url}>.'
                else:  # APA
                    formatted = f'{author}. ({year}). *{title}*. {url}'

                output_lines.append(f"* {formatted}")
                continue

        output_lines.append(line)

    # If no references section existed, but links exist in the document, append one
    if not has_existing_section:
        found_links = []
        seen = set()
        for l in lines:
            for lm in re.finditer(r"\[([^\]]+)\]\((https?://[^\s)]+)\)", l):
                lt, lu = lm.group(1).strip(), lm.group(2).strip()
                if lu not in seen:
                    seen.add(lu)
                    found_links.append((lt, lu))

        if found_links:
            output_lines.append(f"\n\n## References\n\n*Citations formatted in {style.upper()} style.*\n")
            for lt, lu in found_links:
                dom = urlparse(lu).netloc.replace("www.", "")
                yr_m = re.search(r"\b(19\d\d|20\d\d)\b", lt)
                yr = yr_m.group(1) if yr_m else "n.d."
                src = re.sub(r"\s*\(\d{4}\)", "", lt).strip() or dom
                if style == "mla":
                    output_lines.append(f'* "{src}." *{dom}*, {yr}, {lu}.')
                elif style == "harvard":
                    output_lines.append(f'* {src} ({yr}) *{src}*. Available at: <{lu}>.')
                else:
                    output_lines.append(f'* {src}. ({yr}). {lu}')

    return "\n".join(output_lines)


# ─── Table Parsing Helper ───────────────────────────────────────────────────

def _parse_markdown_table(lines: List[str]) -> Tuple[List[str], List[List[str]]]:
    """
    Parse consecutive markdown table lines into (headers, rows).
    """
    headers: List[str] = []
    rows: List[List[str]] = []
    if not lines:
        return headers, rows

    def split_row(row_str: str) -> List[str]:
        s = row_str.strip()
        if s.startswith("|"):
            s = s[1:]
        if s.endswith("|"):
            s = s[:-1]
        return [c.strip() for c in s.split("|")]

    first_cells = split_row(lines[0])
    if len(lines) > 1:
        second_cells = split_row(lines[1])
        is_divider = all(re.match(r"^:?-+:?$", c) for c in second_cells if c)
        if is_divider:
            headers = first_cells
            for l in lines[2:]:
                rows.append(split_row(l))
            return headers, rows

    for l in lines:
        rows.append(split_row(l))
    return headers, rows


# ─── DOCX Export ─────────────────────────────────────────────────────────────

def _export_docx(markdown_text: str, output_path: Path) -> Path:
    """
    Convert *markdown_text* to a styled .docx file at *output_path*.
    """
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()

    # ── Page margins ──────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # ── Custom style helpers ──────────────────────────────────────
    def _add_heading(text: str, level: int) -> None:
        doc.add_heading(text.strip(), level=level)

    def _add_paragraph(text: str, style: str = "Normal") -> None:
        p = doc.add_paragraph(style=style)
        parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("`") and part.endswith("`"):
                run = p.add_run(part[1:-1])
                run.font.name = "Courier New"
                run.font.size = Pt(9)
            elif part.startswith("*") and part.endswith("*"):
                run = p.add_run(part[1:-1])
                run.italic = True
            else:
                p.add_run(part)

    def _add_code_block(code: str) -> None:
        p = doc.add_paragraph(style="Normal")
        run = p.add_run(code)
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F0F0F0")
        pPr.append(shd)

    def _add_horizontal_rule() -> None:
        doc.add_paragraph("─" * 60, style="Normal")

    def _render_frontmatter(fm_text: str) -> None:
        doc.add_heading("Document Metadata", level=2)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Shading Accent 1"
        for line in fm_text.strip().splitlines():
            if ":" in line:
                key, _, val = line.partition(":")
                row = table.add_row()
                row.cells[0].text = key.strip()
                row.cells[1].text = val.strip().strip('"')
        doc.add_paragraph()

    def _add_docx_table(headers: List[str], rows: List[List[str]]) -> None:
        num_cols = max(len(headers), max((len(r) for r in rows), default=0))
        if num_cols == 0:
            return
        table = doc.add_table(rows=0, cols=num_cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        if headers:
            hdr_row = table.add_row()
            trPr = hdr_row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:tblHeader"))
            for j in range(num_cols):
                h_text = headers[j] if j < len(headers) else ""
                cell = hdr_row.cells[j]
                cell.text = h_text
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "1A2F5E")
                cell._tc.get_or_add_tcPr().append(shd)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(9.5)

        for row_idx, r_data in enumerate(rows):
            row = table.add_row()
            is_even = (row_idx % 2 == 1)
            for j in range(num_cols):
                val = r_data[j] if j < len(r_data) else ""
                cell = row.cells[j]
                cell.text = val
                if is_even:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "F7F8FA")
                    cell._tc.get_or_add_tcPr().append(shd)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        doc.add_paragraph()

    # ── Parse and render ──────────────────────────────────────────
    lines = markdown_text.splitlines()
    i = 0
    in_code_block = False
    code_lines: List[str] = []
    frontmatter_lines: List[str] = []
    in_frontmatter = False

    while i < len(lines):
        line = lines[i]

        # YAML frontmatter
        if i == 0 and line.strip() == "---":
            in_frontmatter = True
            i += 1
            continue
        if in_frontmatter:
            if line.strip() == "---":
                in_frontmatter = False
                _render_frontmatter("\n".join(frontmatter_lines))
            else:
                frontmatter_lines.append(line)
            i += 1
            continue

        # Fenced code blocks
        if line.strip().startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                _add_code_block("\n".join(code_lines))
                in_code_block = False
                code_lines = []
            i += 1
            continue
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Markdown tables
        if line.strip().startswith("|") and line.strip().endswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            headers, rows = _parse_markdown_table(tbl_lines)
            _add_docx_table(headers, rows)
            continue

        # Headings
        if line.startswith("### "):
            _add_heading(line[4:], 3)
        elif line.startswith("## "):
            _add_heading(line[3:], 2)
        elif line.startswith("# "):
            _add_heading(line[2:], 1)

        # Horizontal rule
        elif re.match(r"^-{3,}$", line.strip()) or re.match(r"^={3,}$", line.strip()):
            _add_horizontal_rule()

        # Bullet list
        elif re.match(r"^[\-\*\+]\s+", line):
            _add_paragraph(re.sub(r"^[\-\*\+]\s+", "", line), style="List Bullet")

        # Numbered list
        elif re.match(r"^\d+\.\s+", line):
            _add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number")

        # Blockquote
        elif line.startswith("> "):
            _add_paragraph(line[2:], style="Normal")

        # Blank line — skip
        elif line.strip() == "":
            pass

        # Regular paragraph
        else:
            _add_paragraph(line)

        i += 1

    doc.save(str(output_path))
    return output_path


# ─── TXT Export ──────────────────────────────────────────────────────────────

def _export_txt(markdown_text: str, output_path: Path) -> Path:
    """
    Convert *markdown_text* to a well-formatted plain text report at *output_path*.
    Maintains clean ASCII table alignment, structured section headings, and bullet points.
    """
    lines = markdown_text.splitlines()
    out = []
    i = 0
    in_code_block = False
    in_frontmatter = False

    while i < len(lines):
        line = lines[i]
        if i == 0 and line.strip() == "---":
            in_frontmatter = True
            i += 1
            continue
        if in_frontmatter:
            if line.strip() == "---":
                in_frontmatter = False
            else:
                out.append(line)
            i += 1
            continue

        if line.strip().startswith("```"):
            in_code_block = not in_code_block
            out.append("------------------------------------------------------------")
            i += 1
            continue

        if in_code_block:
            out.append("    " + line)
            i += 1
            continue

        # Headings
        if line.startswith("# "):
            title = line[2:].strip().upper()
            out.append("\n" + "=" * len(title))
            out.append(title)
            out.append("=" * len(title) + "\n")
            i += 1
            continue
        elif line.startswith("## "):
            title = line[3:].strip()
            out.append("\n" + "-" * len(title))
            out.append(title)
            out.append("-" * len(title) + "\n")
            i += 1
            continue
        elif line.startswith("### "):
            title = line[4:].strip()
            out.append("\n[ " + title + " ]\n")
            i += 1
            continue

        # Tables
        if line.strip().startswith("|") and line.strip().endswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            headers, rows = _parse_markdown_table(tbl_lines)
            if headers or rows:
                all_rows = ([headers] if headers else []) + rows
                num_cols = max(len(r) for r in all_rows)
                col_widths = [0] * num_cols
                for r in all_rows:
                    for c_idx in range(num_cols):
                        val = r[c_idx] if c_idx < len(r) else ""
                        col_widths[c_idx] = max(col_widths[c_idx], min(len(val), 40))
                # Build ASCII table
                sep = "+" + "+".join("-" * (w + 2) for w in col_widths) + "+"
                out.append(sep)
                if headers:
                    hdr_line = "|" + "|".join(f" {headers[j][:col_widths[j]].ljust(col_widths[j])} " if j < len(headers) else " " * (col_widths[j] + 2) for j in range(num_cols)) + "|"
                    out.append(hdr_line)
                    out.append(sep)
                for r in rows:
                    r_line = "|" + "|".join(f" {r[j][:col_widths[j]].ljust(col_widths[j])} " if j < len(r) else " " * (col_widths[j] + 2) for j in range(num_cols)) + "|"
                    out.append(r_line)
                out.append(sep + "\n")
            continue

        # Bullet list
        if re.match(r"^[\-\*\+]\s+", line):
            content = re.sub(r"^[\-\*\+]\s+", "", line)
            content = re.sub(r"\*\*([^*]+)\*\*", r"\1", content)
            content = re.sub(r"\*([^*]+)\*", r"\1", content)
            out.append(f"  • {content}")
            i += 1
            continue

        # Numbered list
        if re.match(r"^\d+\.\s+", line):
            content = re.sub(r"^\d+\.\s+", "", line)
            m_num = re.match(r"^(\d+)\.", line)
            num = m_num.group(1) if m_num else "1"
            content = re.sub(r"\*\*([^*]+)\*\*", r"\1", content)
            out.append(f"  {num}. {content}")
            i += 1
            continue

        # Blockquote
        if line.startswith("> "):
            out.append("  | " + line[2:])
            i += 1
            continue

        # Regular paragraphs with inline markdown cleaned
        clean_line = re.sub(r"\*\*([^*]+)\*\*", r"\1", line)
        clean_line = re.sub(r"\*([^*]+)\*", r"\1", clean_line)
        clean_line = re.sub(r"`([^`]+)`", r"\1", clean_line)
        clean_line = re.sub(r"\[([^\]]+)\]\((https?://[^\s)]+)\)", r"\1 (\2)", clean_line)
        out.append(clean_line)
        i += 1

    txt_content = "\n".join(out)
    output_path.write_text(txt_content, encoding="utf-8")
    return output_path


# ─── PDF Export (ReportLab fallback with rich tables) ────────────────────────

def _build_pdf_table(headers: List[str], rows: List[List[str]], styles):
    from reportlab.platypus import Table, TableStyle, Paragraph
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import cm

    head_style = ParagraphStyle(
        "TH",
        parent=styles["Normal"],
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        fontName="Helvetica-Bold",
    )
    body_style = ParagraphStyle(
        "TD",
        parent=styles["Normal"],
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#222222"),
        fontName="Helvetica",
    )

    data = []
    if headers:
        hdr_row = []
        for h in headers:
            safe = h.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            safe = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)
            safe = re.sub(r"\*(.+?)\*", r"<i>\1</i>", safe)
            hdr_row.append(Paragraph(safe, head_style))
        data.append(hdr_row)

    for row in rows:
        b_row = []
        for cell in row:
            safe = cell.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            safe = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)
            safe = re.sub(r"\*(.+?)\*", r"<i>\1</i>", safe)
            b_row.append(Paragraph(safe, body_style))
        data.append(b_row)

    if not data:
        return None

    num_cols = max(len(r) for r in data)
    for r in data:
        while len(r) < num_cols:
            r.append(Paragraph("", body_style))

    # A4 printable width: 21cm - 2 * 2.5cm = 16.0cm
    col_width = (16.0 / max(num_cols, 1)) * cm
    t = Table(data, colWidths=[col_width] * num_cols)

    t_style = [
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d0d5dd')),
    ]
    if headers:
        t_style.append(('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a2f5e')))
        start_row = 1
    else:
        start_row = 0

    for r in range(start_row, len(data)):
        if r % 2 == 0:
            t_style.append(('BACKGROUND', (0, r), (-1, r), colors.HexColor('#f7f8fa')))

    t.setStyle(TableStyle(t_style))
    return t


def _export_pdf_reportlab(markdown_text: str, output_path: Path) -> Path:
    """PDF generation via ReportLab with tables, headings, and lists."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Preformatted,
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        "H1Style", parent=styles["Heading1"], fontSize=18, leading=22, textColor=colors.HexColor("#1a2f5e"),
    ))
    styles.add(ParagraphStyle(
        "H2Style", parent=styles["Heading2"], fontSize=14, leading=18, textColor=colors.HexColor("#1a2f5e"),
    ))
    styles.add(ParagraphStyle(
        "H3Style", parent=styles["Heading3"], fontSize=11, leading=15, textColor=colors.HexColor("#2c4a7c"),
    ))

    doc = SimpleDocTemplate(
        str(output_path), pagesize=A4,
        leftMargin=2.5*cm, rightMargin=2.5*cm,
        topMargin=2.5*cm, bottomMargin=2.5*cm,
    )

    story = []
    lines = markdown_text.splitlines()
    i = 0
    in_code_block = False
    code_acc: List[str] = []
    in_frontmatter = False

    while i < len(lines):
        line = lines[i]
        if i == 0 and line.strip() == "---":
            in_frontmatter = True
            i += 1
            continue
        if in_frontmatter:
            if line.strip() == "---":
                in_frontmatter = False
            i += 1
            continue

        if line.strip().startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_acc = []
            else:
                story.append(Preformatted("\n".join(code_acc), styles["Code"]))
                story.append(Spacer(1, 0.2*cm))
                in_code_block = False
                code_acc = []
            i += 1
            continue
        if in_code_block:
            code_acc.append(line)
            i += 1
            continue

        # Markdown tables
        if line.strip().startswith("|") and line.strip().endswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            headers, rows = _parse_markdown_table(tbl_lines)
            tbl = _build_pdf_table(headers, rows, styles)
            if tbl:
                story.append(tbl)
                story.append(Spacer(1, 0.25*cm))
            continue

        if line.startswith("# "):
            story.append(Paragraph(line[2:], styles["H1Style"]))
            story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1a2f5e")))
            story.append(Spacer(1, 0.2*cm))
        elif line.startswith("## "):
            story.append(Spacer(1, 0.25*cm))
            story.append(Paragraph(line[3:], styles["H2Style"]))
            story.append(Spacer(1, 0.15*cm))
        elif line.startswith("### "):
            story.append(Spacer(1, 0.2*cm))
            story.append(Paragraph(line[4:], styles["H3Style"]))
            story.append(Spacer(1, 0.1*cm))
        elif re.match(r"^-{3,}$", line.strip()) or re.match(r"^={3,}$", line.strip()):
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.grey))
            story.append(Spacer(1, 0.15*cm))
        elif re.match(r"^[\-\*]\s+", line):
            safe = re.sub(r"^[\-\*]\s+", "", line)
            safe = safe.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            safe = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)
            safe = re.sub(r"\*(.+?)\*", r"<i>\1</i>", safe)
            txt = "• " + safe
            story.append(Paragraph(txt, styles["Normal"]))
        elif re.match(r"^\d+\.\s+", line):
            safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            safe = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)
            safe = re.sub(r"\*(.+?)\*", r"<i>\1</i>", safe)
            story.append(Paragraph(safe, styles["Normal"]))
        elif line.strip() == "":
            story.append(Spacer(1, 0.2*cm))
        else:
            safe = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            safe = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)
            safe = re.sub(r"\*(.+?)\*", r"<i>\1</i>", safe)
            story.append(Paragraph(safe, styles["Normal"]))

        i += 1

    doc.build(story)
    return output_path


def _export_pdf(markdown_text: str, output_path: Path) -> Tuple[Path, str]:
    """
    Export PDF via ReportLab (or WeasyPrint if available and working).
    """
    try:
        from weasyprint import HTML, CSS  # type: ignore
        import markdown as md  # type: ignore
        body = md.markdown(markdown_text, extensions=["tables", "fenced_code", "toc", "nl2br"])
        html = f"<!DOCTYPE html><html><head><meta charset='utf-8'></head><body>{body}</body></html>"
        HTML(string=html).write_pdf(str(output_path))
        return output_path, "weasyprint"
    except Exception:
        # Fall back reliably to ReportLab
        _export_pdf_reportlab(markdown_text, output_path)
        return output_path, "reportlab"


# ─── Main Exporter Class ─────────────────────────────────────────────────────

class ReportExporter:
    """
    Multi-format report exporter: Markdown → PDF, DOCX, Markdown, TXT.
    """

    def __init__(self, skip_pdf: bool = False, skip_docx: bool = False) -> None:
        self.skip_pdf = skip_pdf
        self.skip_docx = skip_docx

    def export_single(
        self,
        markdown_text: str,
        fmt: str,
        output_path: Path,
        citation_style: str = "apa",
    ) -> Path:
        """
        Export markdown_text to a single requested format (pdf, docx, md, txt),
        applying the selected citation style.
        """
        fmt_clean = (fmt or "pdf").lower().strip().lstrip(".")
        styled_md = apply_citation_style(markdown_text, style=citation_style)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if fmt_clean == "docx":
            return _export_docx(styled_md, output_path.with_suffix(".docx"))
        elif fmt_clean == "pdf":
            saved, _ = _export_pdf(styled_md, output_path.with_suffix(".pdf"))
            return saved
        elif fmt_clean == "txt":
            return _export_txt(styled_md, output_path.with_suffix(".txt"))
        elif fmt_clean == "md":
            target = output_path.with_suffix(".md")
            target.write_text(styled_md, encoding="utf-8")
            return target
        else:
            raise ValueError(f"Unsupported export format: {fmt}")

    def export_all(
        self,
        markdown_text: str,
        base_path: Path,
        citation_style: str = "apa",
    ) -> Dict[str, Optional[Path]]:
        """
        Save the report in all supported formats (md, docx, pdf, txt).
        """
        styled_md = apply_citation_style(markdown_text, style=citation_style)
        paths: Dict[str, Optional[Path]] = {}

        # ── MD ───────────────────────────────────────────────────
        md_path = base_path.with_suffix(".md")
        try:
            md_path.write_text(styled_md, encoding="utf-8")
            paths["md"] = md_path.resolve()
        except Exception:
            paths["md"] = None

        # ── TXT ──────────────────────────────────────────────────
        txt_path = base_path.with_suffix(".txt")
        try:
            _export_txt(styled_md, txt_path)
            paths["txt"] = txt_path.resolve()
        except Exception:
            paths["txt"] = None

        # ── DOCX ─────────────────────────────────────────────────
        if not self.skip_docx:
            docx_path = base_path.with_suffix(".docx")
            try:
                _export_docx(styled_md, docx_path)
                paths["docx"] = docx_path.resolve()
            except Exception as exc:
                console.print(f"  [dim red]⚠ DOCX export failed: {exc}[/dim red]")
                paths["docx"] = None
        else:
            paths["docx"] = None

        # ── PDF ──────────────────────────────────────────────────
        if not self.skip_pdf:
            pdf_path = base_path.with_suffix(".pdf")
            try:
                saved_path, engine = _export_pdf(styled_md, pdf_path)
                paths["pdf"] = saved_path.resolve()
            except Exception as exc:
                console.print(f"  [dim red]⚠ PDF export failed: {exc}[/dim red]")
                paths["pdf"] = None
        else:
            paths["pdf"] = None

        return paths
